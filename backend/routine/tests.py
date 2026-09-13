import io
from datetime import time

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from . import imports as bulk_imports
from . import scheduling
from .exporters import export_routines_excel, export_routines_pdf
from .models import Course, Room, Routine, RoutineRequirement, Teacher, TimeSlot


def make_teacher(username, email):
    user = User.objects.create_user(username=username, password='pass1234', email=email, first_name=username)
    return Teacher.objects.create(user=user, name=username, email=email)


class SchedulingConflictTests(TestCase):
    def setUp(self):
        self.teacher = make_teacher('t1', 't1@example.com')
        self.course = Course.objects.create(name='Data Structures', code='CSE1101')
        self.room = Room.objects.create(name='401(MB)', room_type=Room.THEORY)
        self.slot = TimeSlot.objects.create(start_time=time(9, 0), end_time=time(10, 30), order=1)

    def test_no_conflict_for_first_booking(self):
        scheduling.assert_no_conflicts(day='Saturday', time_slot=self.slot, teacher=self.teacher, room=self.room, section='1A')

    def test_teacher_double_booking_raises(self):
        Routine.objects.create(
            teacher=self.teacher, course=self.course, day='Saturday', section='1A',
            room_ref=self.room, time_slot=self.slot,
            start_time=self.slot.start_time, end_time=self.slot.end_time, room=self.room.name,
        )
        other_room = Room.objects.create(name='402(MB)', room_type=Room.THEORY)
        with self.assertRaises(scheduling.ConflictError):
            scheduling.assert_no_conflicts(day='Saturday', time_slot=self.slot, teacher=self.teacher, room=other_room, section='1B')

    def test_room_double_booking_raises(self):
        Routine.objects.create(
            teacher=self.teacher, course=self.course, day='Saturday', section='1A',
            room_ref=self.room, time_slot=self.slot,
            start_time=self.slot.start_time, end_time=self.slot.end_time, room=self.room.name,
        )
        other_teacher = make_teacher('t2', 't2@example.com')
        with self.assertRaises(scheduling.ConflictError):
            scheduling.assert_no_conflicts(day='Saturday', time_slot=self.slot, teacher=other_teacher, room=self.room, section='1B')

    def test_excludes_own_routine_id_on_update(self):
        routine = Routine.objects.create(
            teacher=self.teacher, course=self.course, day='Saturday', section='1A',
            room_ref=self.room, time_slot=self.slot,
            start_time=self.slot.start_time, end_time=self.slot.end_time, room=self.room.name,
        )
        scheduling.assert_no_conflicts(
            day='Saturday', time_slot=self.slot, teacher=self.teacher, room=self.room,
            section='1A', exclude_routine_id=routine.id,
        )


class GeneratorTests(TestCase):
    def setUp(self):
        self.teacher = make_teacher('gen1', 'gen1@example.com')
        self.course = Course.objects.create(name='Algorithms', code='CSE2101')
        Room.objects.create(name='R1', room_type=Room.THEORY)
        TimeSlot.objects.create(start_time=time(9, 0), end_time=time(10, 30), order=1)
        TimeSlot.objects.create(start_time=time(10, 40), end_time=time(12, 10), order=2)

    def test_places_all_sessions_when_capacity_available(self):
        requirement = RoutineRequirement.objects.create(
            course=self.course, teacher=self.teacher, section='2A', sessions_per_week=2,
        )
        result = scheduling.generate_preview(requirement_ids=[requirement.id])
        self.assertEqual(result['summary']['placed_count'], 2)
        self.assertEqual(result['summary']['unplaced_count'], 0)

    def test_reports_unplaced_when_no_room_of_required_type(self):
        requirement = RoutineRequirement.objects.create(
            course=self.course, teacher=self.teacher, section='2A', sessions_per_week=1,
            required_room_type=Room.LAB,
        )
        result = scheduling.generate_preview(requirement_ids=[requirement.id])
        self.assertEqual(result['summary']['placed_count'], 0)
        self.assertEqual(result['summary']['unplaced_count'], 1)
        self.assertIn('reason', result['unplaced'][0])


class SpecializationGeneratorTests(TestCase):
    def setUp(self):
        self.course = Course.objects.create(name='Math II', code='MATH1302')
        Room.objects.create(name='R1', room_type=Room.THEORY)
        TimeSlot.objects.create(start_time=time(9, 0), end_time=time(10, 30), order=1)
        TimeSlot.objects.create(start_time=time(10, 40), end_time=time(12, 10), order=2)

    def test_auto_assigns_sole_qualified_teacher(self):
        teacher = make_teacher('spec1', 'spec1@example.com')
        teacher.specializations.add(self.course)
        requirement = RoutineRequirement.objects.create(
            course=self.course, teacher=None, section='2A', sessions_per_week=1,
        )
        result = scheduling.generate_preview(requirement_ids=[requirement.id])
        self.assertEqual(result['summary']['placed_count'], 1)
        self.assertEqual(result['placed'][0]['teacher_id'], teacher.id)

    def test_reports_unplaced_when_no_teacher_specialized(self):
        requirement = RoutineRequirement.objects.create(
            course=self.course, teacher=None, section='2A', sessions_per_week=1,
        )
        result = scheduling.generate_preview(requirement_ids=[requirement.id])
        self.assertEqual(result['summary']['unplaced_count'], 1)
        self.assertIn('specialized', result['unplaced'][0]['reason'])

    def test_load_balances_across_qualified_teachers(self):
        t1 = make_teacher('spec2', 'spec2@example.com')
        t2 = make_teacher('spec3', 'spec3@example.com')
        t1.specializations.add(self.course)
        t2.specializations.add(self.course)
        requirement = RoutineRequirement.objects.create(
            course=self.course, teacher=None, section='2A', sessions_per_week=2,
        )
        result = scheduling.generate_preview(requirement_ids=[requirement.id])
        self.assertEqual(result['summary']['placed_count'], 2)
        assigned_teacher_ids = {row['teacher_id'] for row in result['placed']}
        self.assertEqual(assigned_teacher_ids, {t1.id, t2.id})


def _csv_file(text):
    return io.BytesIO(text.encode('utf-8'))


class CsvImportTests(TestCase):
    def test_import_courses_creates_rows_and_reports_duplicates(self):
        csv_text = "name,code,credit\nDatabases,CSE2201,3\nDatabases,CSE2201,3\n"
        report = bulk_imports.import_courses_from_csv(_csv_file(csv_text))
        self.assertEqual(report['created_count'], 1)
        self.assertEqual(report['failed_count'], 1)
        self.assertEqual(Course.objects.get(code='CSE2201').credit, 3)

    def test_import_rooms_validates_room_type(self):
        csv_text = "name,room_type\n501,Theory\n502,Bogus\n"
        report = bulk_imports.import_rooms_from_csv(_csv_file(csv_text))
        self.assertEqual(report['created_count'], 1)
        self.assertEqual(report['failed_count'], 1)
        self.assertTrue(Room.objects.filter(name='501').exists())

    def test_import_timeslots_validates_time_order(self):
        csv_text = "label,start_time,end_time,order\nP1,09:00,10:30,1\nBad,10:30,09:00,2\n"
        report = bulk_imports.import_timeslots_from_csv(_csv_file(csv_text))
        self.assertEqual(report['created_count'], 1)
        self.assertEqual(report['failed_count'], 1)

    def test_import_teachers_creates_user_and_specializations(self):
        course = Course.objects.create(name='Math', code='MATH1302')
        csv_text = (
            "name,email,username,password,acronym,designation,department,mobile_number,specialization\n"
            "Jane Doe,jane@example.com,janed,StrongPass123!,JD,Lecturer,CSE,12345,MATH1302\n"
        )
        report = bulk_imports.import_teachers_from_csv(_csv_file(csv_text))
        self.assertEqual(report['created_count'], 1)
        self.assertEqual(report['failed_count'], 0)
        teacher = Teacher.objects.get(email='jane@example.com')
        self.assertEqual(teacher.acronym, 'JD')
        self.assertIn(course, teacher.specializations.all())
        self.assertTrue(User.objects.filter(username='janed').exists())

    def test_import_teachers_warns_but_still_creates_on_unknown_specialization_code(self):
        csv_text = (
            "name,email,username,password,acronym,designation,department,mobile_number,specialization\n"
            "Jane Doe,jane@example.com,janed,StrongPass123!,JD,Lecturer,CSE,12345,NOPE9999\n"
        )
        report = bulk_imports.import_teachers_from_csv(_csv_file(csv_text))
        self.assertEqual(report['created_count'], 1)
        self.assertEqual(report['failed_count'], 0)
        self.assertEqual(len(report['warnings']), 1)
        self.assertIn('NOPE9999', report['warnings'][0]['warning'])
        self.assertTrue(User.objects.filter(username='janed').exists())
        self.assertEqual(Teacher.objects.get(email='jane@example.com').specializations.count(), 0)

    def test_import_teachers_matches_specialization_case_insensitively_and_comma_separated(self):
        course_a = Course.objects.create(name='Math', code='MATH1302')
        course_b = Course.objects.create(name='Physics', code='PHY1201')
        csv_text = (
            "name,email,username,password,acronym,designation,department,mobile_number,specialization\n"
            'Jane Doe,jane@example.com,janed,StrongPass123!,JD,Lecturer,CSE,12345,"math1302, phy1201"\n'
        )
        report = bulk_imports.import_teachers_from_csv(_csv_file(csv_text))
        self.assertEqual(report['created_count'], 1)
        self.assertEqual(report['warnings'], [])
        teacher = Teacher.objects.get(email='jane@example.com')
        self.assertEqual(set(teacher.specializations.all()), {course_a, course_b})


class ExportImportTests(TestCase):
    def setUp(self):
        self.teacher = make_teacher('exp1', 'exp1@example.com')
        self.course = Course.objects.create(name='Databases', code='CSE2201')
        self.room = Room.objects.create(name='403(MB)', room_type=Room.THEORY)
        self.slot = TimeSlot.objects.create(start_time=time(9, 0), end_time=time(10, 30), order=1)
        self.routine = Routine.objects.create(
            teacher=self.teacher, course=self.course, day='Saturday', section='3B',
            room_ref=self.room, time_slot=self.slot,
            start_time=self.slot.start_time, end_time=self.slot.end_time, room=self.room.name,
        )

    def test_pdf_export_produces_bytes(self):
        content = export_routines_pdf(Routine.objects.all())
        self.assertTrue(content.startswith(b'%PDF'))

    def test_excel_export_round_trips_through_import(self):
        self.teacher.acronym = 'EX1'
        self.teacher.save()
        content = export_routines_excel(Routine.objects.all())

        import io
        workbook = load_workbook(io.BytesIO(content))
        report = scheduling.import_routine_grid_from_workbook(workbook)
        self.assertEqual(report['created_count'], 1)
        self.assertEqual(report['failed_count'], 0)
        self.assertEqual(Routine.objects.count(), 1)

    def test_import_replaces_existing_routine(self):
        self.teacher.acronym = 'EX1'
        self.teacher.save()
        content = export_routines_excel(Routine.objects.all())

        other_room = Room.objects.create(name='999(MB)', room_type=Room.THEORY)
        Routine.objects.create(
            teacher=self.teacher, course=self.course, day='Sunday', section='9Z',
            room_ref=other_room, time_slot=self.slot,
            start_time=self.slot.start_time, end_time=self.slot.end_time, room=other_room.name,
        )
        self.assertEqual(Routine.objects.count(), 2)

        import io
        workbook = load_workbook(io.BytesIO(content))
        report = scheduling.import_routine_grid_from_workbook(workbook)
        self.assertEqual(report['created_count'], 1)
        self.assertEqual(Routine.objects.count(), 1)
        self.assertFalse(Routine.objects.filter(day='Sunday').exists())

    def test_import_rejects_unknown_room_and_makes_no_changes(self):
        self.teacher.acronym = 'EX1'
        self.teacher.save()
        content = export_routines_excel(Routine.objects.all())

        import io
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook.active
        sheet['A4'] = 'Nonexistent Room'

        report = scheduling.import_routine_grid_from_workbook(workbook)
        self.assertEqual(report['created_count'], 0)
        self.assertGreaterEqual(report['failed_count'], 1)
        self.assertIn('unknown room', report['errors'][0]['error'])
        self.assertEqual(Routine.objects.count(), 1)

    def test_import_rejects_cross_room_teacher_double_booking(self):
        from .exporters import _timeslot_label_ampm

        self.teacher.acronym = 'EX1'
        self.teacher.save()
        other_room = Room.objects.create(name='999(MB)', room_type=Room.THEORY)
        content = export_routines_excel(Routine.objects.all())

        import io
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook.active

        def find_cell(value):
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value == value:
                        return cell.row, cell.column
            return None, None

        day_row, _ = find_cell('Saturday')
        self.assertIsNotNone(day_row)
        header_row = day_row + 1
        _, slot_col = find_cell(_timeslot_label_ampm(self.slot))
        self.assertIsNotNone(slot_col)

        room_row = None
        for candidate_row in range(header_row + 2, sheet.max_row + 1):
            if sheet.cell(row=candidate_row, column=1).value == other_room.name:
                room_row = candidate_row
                break
        self.assertIsNotNone(room_row)

        # Same teacher, a different room, same day+slot as the existing booking.
        sheet.cell(row=room_row, column=slot_col, value=self.course.code)
        sheet.cell(row=room_row, column=slot_col + 1, value='ZZ')
        sheet.cell(row=room_row, column=slot_col + 2, value='EX1')

        report = scheduling.import_routine_grid_from_workbook(workbook)
        self.assertEqual(report['created_count'], 0)
        self.assertEqual(Routine.objects.count(), 1)
        self.assertTrue(any('double-booked' in err['error'] for err in report['errors']))


class RoutineApiPermissionTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin = User.objects.create_superuser(username='admin', password='pass1234', email='admin@example.com')
        self.student = User.objects.create_user(username='student', password='pass1234', email='student@example.com')

    def test_non_admin_cannot_generate(self):
        self.client.force_authenticate(self.student)
        response = self.client.post('/api/routines/generate/', {}, format='json')
        self.assertEqual(response.status_code, 403)

    def test_non_admin_cannot_import_courses_csv(self):
        self.client.force_authenticate(self.student)
        response = self.client.post('/api/courses/import-csv/', {}, format='multipart')
        self.assertEqual(response.status_code, 403)

    def test_admin_can_import_courses_csv(self):
        self.client.force_authenticate(self.admin)
        csv_content = b"name,code,credit\nDatabases,CSE2201,3\n"
        upload = io.BytesIO(csv_content)
        upload.name = 'courses.csv'
        response = self.client.post('/api/courses/import-csv/', {'file': upload}, format='multipart')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['created_count'], 1)

    def test_admin_can_generate_empty_preview(self):
        self.client.force_authenticate(self.admin)
        response = self.client.post('/api/routines/generate/', {}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['summary']['total_sessions'], 0)

    def test_export_pdf_via_api(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get('/api/routines/export/?type=pdf')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def test_export_excel_via_api(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get('/api/routines/export/?type=excel')
        self.assertEqual(response.status_code, 200)

    def test_export_missing_type_returns_400(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get('/api/routines/export/')
        self.assertEqual(response.status_code, 400)

    def test_non_admin_cannot_import_excel(self):
        self.client.force_authenticate(self.student)
        response = self.client.post('/api/routines/import-excel/', {}, format='multipart')
        self.assertEqual(response.status_code, 403)
