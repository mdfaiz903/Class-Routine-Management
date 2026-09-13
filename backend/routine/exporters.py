import io

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Room, Routine, TimeSlot


def _timeslot_label(time_slot):
    return f"{time_slot.start_time.strftime('%H:%M')}-{time_slot.end_time.strftime('%H:%M')}"


def _timeslot_label_ampm(time_slot):
    start = time_slot.start_time.strftime('%I:%M %p').lstrip('0')
    end = time_slot.end_time.strftime('%I:%M %p').lstrip('0')
    return f"{start} - {end}"


def _faculty_label(teacher):
    if not teacher:
        return 'TBA'
    return teacher.acronym or teacher.name


def export_routines_pdf(routines):
    rooms = list(Room.objects.order_by('name'))
    slots = list(TimeSlot.objects.order_by('order', 'start_time'))

    days_order = ['Saturday', 'Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    days_present = [day for day in days_order if any(r.day == day for r in routines)]
    if not days_present:
        days_present = days_order

    by_day_room_slot = {}
    for routine in routines:
        if routine.room_ref_id is None or routine.time_slot_id is None:
            continue
        by_day_room_slot[(routine.day, routine.room_ref_id, routine.time_slot_id)] = routine

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    elements = []

    for index, day in enumerate(days_present):
        if index > 0:
            elements.append(PageBreak())
        elements.append(Paragraph(day, styles['Heading2']))
        elements.append(Spacer(1, 6))

        header_row = ['Room'] + [_timeslot_label(slot) for slot in slots]
        data = [header_row]

        for room in rooms:
            row = [room.name]
            for slot in slots:
                routine = by_day_room_slot.get((day, room.id, slot.id))
                if routine:
                    teacher_label = routine.teacher.name if routine.teacher else 'TBA'
                    cell = f"{routine.course.code} ({routine.section})\n{teacher_label}"
                else:
                    cell = ''
                row.append(cell)
            data.append(row)

        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6366f1')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 18))

    doc.build(elements)
    return buffer.getvalue()


def export_routines_excel(routines):
    """Build the Central_Class_Routine-style grid: day blocks, one row per Room, one
    3-column (Course Code/Section/Faculty) group per TimeSlot. Mirrors what
    scheduling.import_routine_grid_from_workbook expects, so export -> hand-edit ->
    import round-trips.
    """
    rooms = list(Room.objects.order_by('name'))
    slots = list(TimeSlot.objects.order_by('order', 'start_time'))
    days_order = [choice[0] for choice in Routine.DAYS_OF_WEEK]

    by_day_room_slot = {}
    for routine in routines:
        if routine.room_ref_id is None or routine.time_slot_id is None:
            continue
        by_day_room_slot[(routine.day, routine.room_ref_id, routine.time_slot_id)] = routine

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Central_Class_Routine'

    last_col = 1 + 3 * len(slots)
    row = 1
    for day in days_order:
        sheet.cell(row=row, column=2, value=day)
        if last_col > 2:
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
        row += 1

        sheet.cell(row=row, column=1, value='Room')
        col = 2
        for slot in slots:
            sheet.cell(row=row, column=col, value=_timeslot_label_ampm(slot))
            sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
            col += 3
        row += 1

        col = 2
        for _ in slots:
            sheet.cell(row=row, column=col, value='Course Code')
            sheet.cell(row=row, column=col + 1, value='Section')
            sheet.cell(row=row, column=col + 2, value='Faculty')
            col += 3
        row += 1

        for room in rooms:
            sheet.cell(row=row, column=1, value=room.name)
            col = 2
            for slot in slots:
                routine = by_day_room_slot.get((day, room.id, slot.id))
                if routine:
                    sheet.cell(row=row, column=col, value=routine.course.code)
                    sheet.cell(row=row, column=col + 1, value=routine.section)
                    sheet.cell(row=row, column=col + 2, value=_faculty_label(routine.teacher))
                col += 3
            row += 1

        row += 1  # blank separator row between day blocks

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
